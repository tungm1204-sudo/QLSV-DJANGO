"""
Module HR Services
Chứa business logic tạo User và Profile (Student/Lecturer/Staff) một cách nguyên tử.
"""
from django.db import transaction
from rest_framework.exceptions import ValidationError
from .models import Student, Lecturer, Staff, StudentCertificate
from apps.identity.services import UserService
from apps.core.services import log_audit
from apps.identity.models import Role
from apps.master_data.models import Major, AdministrativeClass, EducationSystem, PriorityCategory, Department
import openpyxl
import logging

logger = logging.getLogger(__name__)

class StudentService:
    @staticmethod
    @transaction.atomic
    def create_student(validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> Student:
        email = validated_data.pop('email', None)
        password = validated_data.pop('password', None)
        full_name = validated_data.pop('full_name', None)

        if not email or not password or not full_name:
            raise ValidationError("email, password và full_name là bắt buộc để tạo tài khoản sinh viên.")
        
        # Get role
        try:
            student_role = Role.objects.get(name='Sinh viên')
        except Role.DoesNotExist:
            raise ValidationError("Vai trò 'Sinh viên' chưa được cấu hình trong hệ thống.")

        user_data = {
            'email': email,
            'password': password,
            'full_name': full_name,
            'role': student_role,
            'status': 'ACTIVE',
            'is_staff': False
        }

        # Tạo User Identity
        user = UserService.create_user(user_data, actor_id, ip_address, user_agent)
        
        # Gán user vào profile data
        validated_data['user'] = user

        student = Student.objects.create(**validated_data)
        log_audit(actor_id, 'CREATE', 'Student', {'id': str(student.id), 'student_code': student.student_code}, ip_address, user_agent)
        return student

    @staticmethod
    @transaction.atomic
    def update_student(student: Student, validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> Student:
        # Nếu có thông tin email, full_name, password thì update Identity User
        email = validated_data.pop('email', None)
        password = validated_data.pop('password', None)
        full_name = validated_data.pop('full_name', None)

        if email or password or full_name:
            user_data = {}
            if email: user_data['email'] = email
            if password: user_data['password'] = password
            if full_name: user_data['full_name'] = full_name
            UserService.update_user(student.user, user_data, actor_id, ip_address, user_agent)

        for key, value in validated_data.items():
            setattr(student, key, value)
        student.save()
        log_audit(actor_id, 'UPDATE', 'Student', {'id': str(student.id), 'student_code': student.student_code}, ip_address, user_agent)
        return student

    @staticmethod
    @transaction.atomic
    def delete_student(student: Student, actor_id: str, ip_address: str = None, user_agent: str = None) -> None:
        student.status = 'DROPPED_OUT'
        student.save(update_fields=['status'])
        
        # Khóa tài khoản
        user = student.user
        user.status = 'LOCKED'
        user.save(update_fields=['status'])

        log_audit(actor_id, 'DELETE', 'Student', {'id': str(student.id), 'student_code': student.student_code}, ip_address, user_agent)

    @staticmethod
    def generate_id_card(student: Student, actor_id: str, ip_address: str = None, user_agent: str = None):
        """
        Sinh nội dung thẻ sinh viên (HTML/PDF mock)
        """
        html_content = f"""
        <html>
        <head><title>Thẻ sinh viên - {student.student_code}</title></head>
        <body style="font-family: Arial, sans-serif; text-align: center; border: 1px solid #000; width: 300px; padding: 20px;">
            <h2>ĐẠI HỌC QLSV</h2>
            <h3>THẺ SINH VIÊN</h3>
            <p><strong>Họ tên:</strong> {student.user.full_name}</p>
            <p><strong>MSSV:</strong> {student.student_code}</p>
            <p><strong>Ngày sinh:</strong> {student.date_of_birth.strftime('%d/%m/%Y') if student.date_of_birth else 'N/A'}</p>
            <p><strong>Giới tính:</strong> {dict(student._meta.get_field('gender').choices).get(student.gender, 'N/A')}</p>
            <p><strong>Khóa:</strong> {student.cohort.code if student.cohort else 'N/A'}</p>
            <p><strong>Lớp:</strong> {student.administrative_class.name if student.administrative_class else 'N/A'}</p>
            <p><strong>Ngành:</strong> {student.major.name if student.major else 'N/A'}</p>
            <p><strong>Ngày nhập học:</strong> {student.enrollment_date.strftime('%d/%m/%Y') if student.enrollment_date else 'N/A'}</p>
        </body>
        </html>
        """
        log_audit(actor_id, 'PRINT_ID_CARD', 'Student', {'student_code': student.student_code}, ip_address, user_agent)
        return html_content


class LecturerService:
    @staticmethod
    @transaction.atomic
    def create_lecturer(validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> Lecturer:
        email = validated_data.pop('email', None)
        password = validated_data.pop('password', None)
        full_name = validated_data.pop('full_name', None)

        if not email or not password or not full_name:
            raise ValidationError("email, password và full_name là bắt buộc.")
        
        try:
            role = Role.objects.get(name='Giảng viên')
        except Role.DoesNotExist:
            raise ValidationError("Vai trò 'Giảng viên' chưa được cấu hình.")

        user_data = {
            'email': email,
            'password': password,
            'full_name': full_name,
            'role': role,
            'status': 'ACTIVE',
            'is_staff': True
        }

        user = UserService.create_user(user_data, actor_id, ip_address, user_agent)
        validated_data['user'] = user

        lecturer = Lecturer.objects.create(**validated_data)
        log_audit(actor_id, 'CREATE', 'Lecturer', {'id': str(lecturer.id), 'lecturer_code': lecturer.lecturer_code}, ip_address, user_agent)
        return lecturer

    @staticmethod
    @transaction.atomic
    def update_lecturer(lecturer: Lecturer, validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> Lecturer:
        email = validated_data.pop('email', None)
        password = validated_data.pop('password', None)
        full_name = validated_data.pop('full_name', None)

        if email or password or full_name:
            user_data = {}
            if email: user_data['email'] = email
            if password: user_data['password'] = password
            if full_name: user_data['full_name'] = full_name
            UserService.update_user(lecturer.user, user_data, actor_id, ip_address, user_agent)

        for key, value in validated_data.items():
            setattr(lecturer, key, value)
        lecturer.save()
        log_audit(actor_id, 'UPDATE', 'Lecturer', {'id': str(lecturer.id), 'lecturer_code': lecturer.lecturer_code}, ip_address, user_agent)
        return lecturer

    @staticmethod
    @transaction.atomic
    def delete_lecturer(lecturer: Lecturer, actor_id: str, ip_address: str = None, user_agent: str = None) -> None:
        lecturer.status = 'RESIGNED'
        lecturer.save(update_fields=['status'])
        
        user = lecturer.user
        user.status = 'LOCKED'
        user.save(update_fields=['status'])

        log_audit(actor_id, 'DELETE', 'Lecturer', {'id': str(lecturer.id), 'lecturer_code': lecturer.lecturer_code}, ip_address, user_agent)


class StaffService:
    @staticmethod
    @transaction.atomic
    def create_staff(validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> Staff:
        email = validated_data.pop('email', None)
        password = validated_data.pop('password', None)
        full_name = validated_data.pop('full_name', None)

        if not email or not password or not full_name:
            raise ValidationError("email, password và full_name là bắt buộc.")
        
        role_name = validated_data.pop('role_name', 'Giáo vụ')
        try:
            role = Role.objects.get(name=role_name)
        except Role.DoesNotExist:
            raise ValidationError(f"Vai trò '{role_name}' chưa được cấu hình.")

        user_data = {
            'email': email,
            'password': password,
            'full_name': full_name,
            'role': role,
            'status': 'ACTIVE',
            'is_staff': True
        }

        user = UserService.create_user(user_data, actor_id, ip_address, user_agent)
        validated_data['user'] = user

        staff = Staff.objects.create(**validated_data)
        log_audit(actor_id, 'CREATE', 'Staff', {'id': str(staff.id), 'staff_code': staff.staff_code}, ip_address, user_agent)
        return staff

    @staticmethod
    @transaction.atomic
    def update_staff(staff: Staff, validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> Staff:
        email = validated_data.pop('email', None)
        password = validated_data.pop('password', None)
        full_name = validated_data.pop('full_name', None)

        if email or password or full_name:
            user_data = {}
            if email: user_data['email'] = email
            if password: user_data['password'] = password
            if full_name: user_data['full_name'] = full_name
            UserService.update_user(staff.user, user_data, actor_id, ip_address, user_agent)

        for key, value in validated_data.items():
            setattr(staff, key, value)
        staff.save()
        log_audit(actor_id, 'UPDATE', 'Staff', {'id': str(staff.id), 'staff_code': staff.staff_code}, ip_address, user_agent)
        return staff

    @staticmethod
    @transaction.atomic
    def delete_staff(staff: Staff, actor_id: str, ip_address: str = None, user_agent: str = None) -> None:
        staff.status = 'RESIGNED'
        staff.save(update_fields=['status'])
        
        user = staff.user
        user.status = 'LOCKED'
        user.save(update_fields=['status'])

        log_audit(actor_id, 'DELETE', 'Staff', {'id': str(staff.id), 'staff_code': staff.staff_code}, ip_address, user_agent)


class ImportService:
    @staticmethod
    @transaction.atomic
    def import_students_from_excel(file_obj, actor_id: str, ip_address: str = None, user_agent: str = None):
        try:
            wb = openpyxl.load_workbook(file_obj)
            sheet = wb.active
            created_count = 0
            
            # Format Cột: 
            # 0: email, 1: full_name, 2: password, 3: student_code, 4: major_code, 5: class_code, 6: edu_system_code,
            # 7: admission_type_code, 8: cohort_code, 9: priority_code, 10: status, 
            # 11: enrollment_date, 12: date_of_birth, 13: gender, 14: place_of_birth,
            # 15: ethnicity_code, 16: religion_code, 17: nationality_code,
            # 18: contact_phone, 19: address, 20: id_card_number, 21: health_insurance_number
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                
                # Trích xuất dữ liệu, fallback None nếu file ngắn hơn
                def get_val(idx):
                    return row[idx] if len(row) > idx else None

                email = get_val(0)
                full_name = get_val(1)
                password = get_val(2)
                student_code = get_val(3)
                major_code = get_val(4)
                class_code = get_val(5)
                edu_system_code = get_val(6)
                
                admission_type_code = get_val(7)
                cohort_code = get_val(8)
                priority_code = get_val(9)
                status = get_val(10) or 'ACTIVE'
                
                enrollment_date = get_val(11)
                date_of_birth = get_val(12)
                gender = get_val(13)
                place_of_birth = get_val(14)
                
                ethnicity_code = get_val(15)
                religion_code = get_val(16)
                nationality_code = get_val(17)
                
                contact_phone = get_val(18)
                address = get_val(19)
                id_card_number = get_val(20)
                health_insurance_number = get_val(21)

                major = Major.objects.filter(code=major_code).first() if major_code else None
                admin_class = AdministrativeClass.objects.filter(code=class_code).first() if class_code else None
                edu_system = EducationSystem.objects.filter(code=edu_system_code).first() if edu_system_code else None
                admission_type = AdmissionType.objects.filter(code=admission_type_code).first() if admission_type_code else None
                cohort = Cohort.objects.filter(code=cohort_code).first() if cohort_code else None
                priority = PriorityCategory.objects.filter(code=priority_code).first() if priority_code else None
                
                ethnicity = Ethnicity.objects.filter(code=ethnicity_code).first() if ethnicity_code else None
                religion = Religion.objects.filter(code=religion_code).first() if religion_code else None
                nationality = Nationality.objects.filter(code=nationality_code).first() if nationality_code else None

                if not Student.objects.filter(student_code=student_code).exists():
                    student_data = {
                        'email': email,
                        'full_name': full_name,
                        'password': password,
                        'student_code': student_code,
                        'major': major,
                        'administrative_class': admin_class,
                        'education_system': edu_system,
                        'admission_type': admission_type,
                        'cohort': cohort,
                        'priority_category': priority,
                        'status': status,
                        'enrollment_date': enrollment_date,
                        'date_of_birth': date_of_birth,
                        'gender': gender,
                        'place_of_birth': place_of_birth,
                        'ethnicity': ethnicity,
                        'religion': religion,
                        'nationality': nationality,
                        'contact_phone': contact_phone,
                        'address': address,
                        'id_card_number': id_card_number,
                        'health_insurance_number': health_insurance_number,
                    }
                    StudentService.create_student(student_data, actor_id, ip_address, user_agent)
                    created_count += 1

            if created_count > 0:
                log_audit(actor_id, 'IMPORT_EXCEL', 'Student', {'count': created_count}, ip_address, user_agent)
            return created_count, None
        except Exception as e:
            logger.exception("Import Excel Students failed")
            return 0, str(e)

    @staticmethod
    @transaction.atomic
    def import_lecturers_from_excel(file_obj, actor_id: str, ip_address: str = None, user_agent: str = None):
        try:
            wb = openpyxl.load_workbook(file_obj)
            sheet = wb.active
            created_count = 0
            
            # Giả định cột: 
            # 0: email, 1: full_name, 2: password, 3: lecturer_code, 4: department_code
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                
                email, full_name, password, lecturer_code, department_code = row[0:5]
                
                department = Department.objects.filter(code=department_code).first() if department_code else None

                if not Lecturer.objects.filter(lecturer_code=lecturer_code).exists():
                    lecturer_data = {
                        'email': email,
                        'full_name': full_name,
                        'password': password,
                        'lecturer_code': lecturer_code,
                        'department': department,
                    }
                    LecturerService.create_lecturer(lecturer_data, actor_id, ip_address, user_agent)
                    created_count += 1

            if created_count > 0:
                log_audit(actor_id, 'IMPORT_EXCEL', 'Lecturer', {'count': created_count}, ip_address, user_agent)
            return created_count, None
        except Exception as e:
            logger.exception("Import Excel Lecturers failed")
            return 0, str(e)

    @staticmethod
    @transaction.atomic
    def import_staffs_from_excel(file_obj, actor_id: str, ip_address: str = None, user_agent: str = None):
        try:
            wb = openpyxl.load_workbook(file_obj)
            sheet = wb.active
            created_count = 0
            
            # Giả định cột: 
            # 0: email, 1: full_name, 2: password, 3: staff_code, 4: department_code
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                
                email, full_name, password, staff_code, department_code = row[0:5]
                
                department = Department.objects.filter(code=department_code).first() if department_code else None

                if not Staff.objects.filter(staff_code=staff_code).exists():
                    staff_data = {
                        'email': email,
                        'full_name': full_name,
                        'password': password,
                        'staff_code': staff_code,
                        'department': department,
                    }
                    StaffService.create_staff(staff_data, actor_id, ip_address, user_agent)
                    created_count += 1

            if created_count > 0:
                log_audit(actor_id, 'IMPORT_EXCEL', 'Staff', {'count': created_count}, ip_address, user_agent)
            return created_count, None
        except Exception as e:
            logger.exception("Import Excel Staffs failed")
            return 0, str(e)

class StudentCertificateService:
    @staticmethod
    @transaction.atomic
    def create_certificate(validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> StudentCertificate:
        certificate = StudentCertificate.objects.create(**validated_data)
        log_audit(actor_id, 'CREATE', 'StudentCertificate', {'id': str(certificate.id)}, ip_address, user_agent)
        return certificate

    @staticmethod
    @transaction.atomic
    def update_certificate(certificate: StudentCertificate, validated_data: dict, actor_id: str, ip_address: str = None, user_agent: str = None) -> StudentCertificate:
        for key, value in validated_data.items():
            setattr(certificate, key, value)
        certificate.save()
        log_audit(actor_id, 'UPDATE', 'StudentCertificate', {'id': str(certificate.id)}, ip_address, user_agent)
        return certificate

    @staticmethod
    @transaction.atomic
    def delete_certificate(certificate: StudentCertificate, actor_id: str, ip_address: str = None, user_agent: str = None) -> None:
        certificate.delete()
        log_audit(actor_id, 'DELETE', 'StudentCertificate', {'id': str(certificate.id)}, ip_address, user_agent)

import io
from openpyxl import Workbook
from django.db.models import QuerySet

class ExportService:
    @staticmethod
    def export_queryset_to_excel(queryset: QuerySet, columns: list, title: str):
        wb = Workbook()
        ws = wb.active
        ws.title = title

        # Write header
        ws.append([col[0] for col in columns])

        # Write data
        for obj in queryset:
            row = []
            for col in columns:
                field_getter = col[1]
                try:
                    val = field_getter(obj)
                except Exception:
                    val = ''
                row.append(str(val) if val is not None else '')
            ws.append(row)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

