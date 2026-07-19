"""
Module Identity Services
Chứa toàn bộ logic nghiệp vụ (Business Logic) cốt lõi của phân hệ Identity.
Lý do: Giữ cho View luôn mỏng, dễ dàng viết unit test, và đảm bảo mọi thao tác ghi dữ liệu phức tạp đều được xử lý tập trung (như ghi AuditLog sau mỗi action).
"""

from django.contrib.auth import get_user_model
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError as DjangoValidationError
from rest_framework.exceptions import ValidationError as DRFValidationError
from django.utils import timezone
from django.db import transaction
from .models import AuditLog, OTPToken, LoginSession, SystemConfig, Role, Notification
from .selectors import SystemConfigSelector
import logging
import openpyxl
import random

logger = logging.getLogger(__name__)

User = get_user_model()


def log_audit(user_id, action_name: str, module: str, payload: dict = None, ip_address: str = None, user_agent: str = None, record_id: str = None):
    """
    Hàm tiện ích ghi lại vết hệ thống (Audit Log).
    Lý do: Đảm bảo tuân thủ bảo mật, mọi thao tác thay đổi dữ liệu nhạy cảm đều phải gọi hàm này.
    """
    AuditLog.objects.create(
        user_id=user_id,
        action=action_name,
        module=module,
        payload=payload,
        record_id=record_id,
        ip_address=ip_address,
        user_agent=user_agent
    )


class RoleService:
    """
    Xử lý logic nghiệp vụ cho Role (Tạo, Sửa, Xóa).
    Luôn đi kèm với thao tác ghi AuditLog để theo dõi Admin nào thay đổi quyền hạn.
    """
    @staticmethod
    def create_role(validated_data, actor_id, ip_address=None, user_agent=None):
        # Tạo mới Role trong Database bằng cách unpack (**) dictionary validated_data.
        # Lý do: Dữ liệu này đã được kiểm tra tính hợp lệ (validate) ở tầng Serializer, nên đảm bảo an toàn để lưu thẳng vào DB.
        role = Role.objects.create(**validated_data)
        
        # Ghi vết hành động (Audit Log) lại để lưu lại vết Admin nào (actor_id) đã thao tác tạo Role này.
        # Kèm theo payload chứa ID và tên Role để sau này hệ thống dễ dàng truy vết (traceability).
        log_audit(actor_id, 'CREATE', 'Roles', {'id': str(role.id), 'name': role.name}, ip_address, user_agent)
        
        return role

    @staticmethod
    def update_role(role, validated_data, actor_id, ip_address=None, user_agent=None):
        for key, value in validated_data.items():
            setattr(role, key, value)
        role.save()
        log_audit(actor_id, 'UPDATE', 'Roles', {'id': str(role.id), 'name': role.name}, ip_address, user_agent)
        return role

    @staticmethod
    def delete_role(role, actor_id, ip_address=None, user_agent=None):
        role_info = {'id': str(role.id), 'name': role.name}
        role.delete()
        log_audit(actor_id, 'DELETE', 'Roles', role_info, ip_address, user_agent)


class SystemConfigService:
    """
    Xử lý tạo, cập nhật cấu hình hệ thống.
    Bắt buộc phải ghi log để audit do cấu hình ảnh hưởng trực tiếp đến hệ thống.
    """
    @staticmethod
    @transaction.atomic
    def create_config(validated_data, actor_id, ip_address=None, user_agent=None):
        config = SystemConfig.objects.create(**validated_data)
        log_audit(actor_id, 'CREATE', 'SystemConfig', {'key': config.key}, ip_address, user_agent)
        return config

    @staticmethod
    @transaction.atomic
    def update_config(config, validated_data, actor_id, ip_address=None, user_agent=None):
        for key, value in validated_data.items():
            setattr(config, key, value)
        config.save()
        log_audit(actor_id, 'UPDATE', 'SystemConfig', {'key': config.key}, ip_address, user_agent)
        return config


class NotificationService:
    """
    Xử lý logic liên quan đến Thông báo.
    """
    @staticmethod
    def mark_as_read(notification):
        # Đánh dấu trạng thái đã đọc cho đối tượng notification.
        notification.is_read = True
        
        # Gọi hàm save nhưng chỉ chỉ định trường update_fields=['is_read'].
        # Lý do: Tối ưu hoá câu lệnh SQL UPDATE, chỉ cập nhật 1 cột duy nhất thay vì toàn bộ các cột, giúp DB chạy nhanh hơn.
        notification.save(update_fields=['is_read'])


class AuthService:
    """
    Xử lý logic xác thực, khóa tài khoản (Lockout), tạo mã OTP.
    Lý do: Tách biệt logic kiểm tra security (login fail, expired OTP) ra khỏi serializer.
    """
    @staticmethod
    def check_lockout(email: str) -> tuple[bool, str]:
        try:
            # Query tìm User trong bảng identity_users dựa vào email truyền vào.
            user = User.objects.get(email=email)
            
            # Kiểm tra xem tài khoản có đang bị khóa (cột locked_until có giá trị) 
            # VÀ thời gian khóa vẫn lớn hơn thời gian thực tại server (timezone.now()) hay không.
            # Lý do: Nếu đã qua mốc thời gian khóa, hệ thống sẽ bỏ qua logic này và tự động cho phép login tiếp.
            if user.locked_until:
                if user.locked_until > timezone.now():
                    delta = user.locked_until - timezone.now()
                    minutes_left = int(delta.total_seconds() // 60) + 1
                    return True, f"Tài khoản đã bị khóa do nhập sai quá nhiều lần. Vui lòng thử lại sau {minutes_left} phút."
                else:
                    # Fix: Khi thời gian khóa đã hết, phải reset lại số lần sai về 0
                    # để người dùng được phép nhập lại từ đầu (được sai thêm 5 lần nữa)
                    # Nếu không reset, họ chỉ cần nhập sai 1 lần là bị khóa lại ngay lập tức.
                    user.failed_login_attempts = 0
                    user.locked_until = None
                    user.save(update_fields=['failed_login_attempts', 'locked_until'])
            
            # Không bị khóa hoặc đã hết hạn khóa -> Trả về False (cho phép đi tiếp)
            return False, None
        except User.DoesNotExist:
            # Bắt lỗi Exception nếu email truyền vào không tồn tại trong hệ thống.
            # Vẫn trả về False để không làm rò rỉ (leak) thông tin là email có tồn tại hay không cho hacker biết.
            return False, None

    @staticmethod
    def clear_lockout(user):
        # Reset số lần nhập sai về 0 và xóa trạng thái khóa.
        user.failed_login_attempts = 0
        user.locked_until = None
        
        # Chỉ lưu 2 cột bị thay đổi xuống DB.
        user.save(update_fields=['failed_login_attempts', 'locked_until'])

    @staticmethod
    def generate_otp(email, otp_type=OTPToken.TypeChoices.LOGIN):
        try:
            user = User.objects.get(email=email)
            
            # Sinh ra chuỗi 6 số ngẫu nhiên làm mã OTP.
            code = str(random.randint(100000, 999999))
            
            # Lưu mã OTP vào database, thiết lập loại OTP (Login/ResetPass).
            # Lý do thiết lập thời hạn 5 phút: Ngăn chặn hacker có thời gian rảnh rỗi để vét cạn (brute-force) mã OTP.
            OTPToken.objects.create(
                user=user,
                code=code,
                type=otp_type,
                expires_at=timezone.now() + timezone.timedelta(minutes=5)
            )
            return user, code
        except User.DoesNotExist:
            # Ngụy trang bằng cách trả về None nếu email không tồn tại.
            # Lý do: Không trả về lỗi "Email not found" để tránh bị hacker dùng bot quét thu thập danh sách email của người dùng (Email Enumeration).
            return None, None

    @staticmethod
    @transaction.atomic
    def verify_otp(email, code, otp_type=OTPToken.TypeChoices.LOGIN):
        """
        Fix #10: Lọc theo otp_type để OTP đăng nhập không dùng được cho reset password.
        """
        try:
            user = User.objects.get(email=email)
            token = OTPToken.objects.select_for_update().filter(
                user=user, code=code, type=otp_type, is_used=False
            ).order_by('-created_at').first()

            if not token or not token.is_valid():
                return None, "Invalid or expired OTP"

            token.is_used = True
            token.save(update_fields=['is_used'])

            user.failed_login_attempts = 0
            user.locked_until = None
            user.save(update_fields=['failed_login_attempts', 'locked_until'])

            return user, None
        except User.DoesNotExist:
            return None, "Invalid or expired OTP"

    @staticmethod
    @transaction.atomic
    def reset_password(user, new_password):
        user.set_password(new_password)
        user.save()
        return True

    @staticmethod
    @transaction.atomic
    def record_login(user, ip_address=None, user_agent=None, current_refresh_token=None):
        """
        Fix #5: Bọc trong @transaction.atomic — ghi user.last_login và tạo LoginSession là 1 đơn vị.
        """
        user.last_login = timezone.now()
        user.save(update_fields=['last_login'])

        # Invalidate old sessions
        LoginSession.objects.filter(user=user, is_active=True).update(is_active=False)

        # Blacklist old outstanding tokens to prevent concurrent logins
        from rest_framework_simplejwt.token_blacklist.models import OutstandingToken, BlacklistedToken
        outstanding_tokens = OutstandingToken.objects.filter(user=user)
        
        if current_refresh_token:
            outstanding_tokens = outstanding_tokens.exclude(token=current_refresh_token)
            
        for token in outstanding_tokens:
            BlacklistedToken.objects.get_or_create(token=token)

        LoginSession.objects.create(
            user=user,
            ip_address=ip_address,
            user_agent=user_agent,
            device_info="Unknown",
            is_active=True
        )

    @staticmethod
    def handle_failed_login(email: str) -> None:
        try:
            user = User.objects.get(email=email)
            user.failed_login_attempts += 1

            max_attempts, lockout_duration = SystemConfigSelector.get_lockout_config()

            if user.failed_login_attempts >= max_attempts:
                user.locked_until = timezone.now() + timezone.timedelta(minutes=lockout_duration)
            user.save(update_fields=['failed_login_attempts', 'locked_until'])
        except User.DoesNotExist:
            pass

    @staticmethod
    @transaction.atomic
    def revoke_session(session_id: str, user_id: str, ip_address: str, user_agent: str) -> None:
        """
        Thu hồi một phiên đăng nhập cụ thể của người dùng.
        """
        session = LoginSession.objects.filter(user_id=user_id).get(pk=session_id)
        
        session.is_active = False
        session.save(update_fields=['is_active'])
        
        log_audit(user_id, 'REVOKE_SESSION', 'Auth', {'session_id': str(session_id)}, ip_address, user_agent)


class UserService:
    @staticmethod
    @transaction.atomic
    def create_user(validated_data, actor_id, ip_address=None, user_agent=None):
        """
        Fix #6: Bọc trong @transaction.atomic — tạo user và set_password là 1 đơn vị nguyên tử.
        Fix #14: Validate password bằng Django validators.
        """
        password = validated_data.pop('password', None)

        # Validate password strength trước khi tạo user
        if password:
            try:
                validate_password(password)
            except DjangoValidationError as e:
                raise DRFValidationError({'password': list(e.messages)})

        user = User(**validated_data)
        if password:
            user.set_password(password)
        else:
            user.set_unusable_password()
        user.save()

        log_audit(actor_id, 'CREATE', 'Users', {'id': str(user.id), 'email': user.email}, ip_address, user_agent)
        return user

    @staticmethod
    @transaction.atomic
    def update_user(user, validated_data, actor_id, ip_address=None, user_agent=None):
        password = validated_data.pop('password', None)
        
        if validated_data.get('status') == 'LOCKED':
            if user.is_superuser:
                raise DRFValidationError({'status': 'Không thể khóa tài khoản Super Admin.'})
            if user.id == actor_id:
                raise DRFValidationError({'status': 'Không thể tự khóa tài khoản của chính mình.'})

        for key, value in validated_data.items():
            setattr(user, key, value)
            
        if validated_data.get('status') == 'ACTIVE':
            user.failed_login_attempts = 0
            user.locked_until = None

        if password:
            try:
                validate_password(password, user=user)
            except DjangoValidationError as e:
                raise DRFValidationError({'password': list(e.messages)})
            user.set_password(password)
        user.save()
        log_audit(actor_id, 'UPDATE', 'Users', {'id': str(user.id), 'email': user.email}, ip_address, user_agent)
        return user

    @staticmethod
    @transaction.atomic
    def delete_user(user, actor_id, ip_address=None, user_agent=None):
        user_info = {'id': str(user.id), 'email': user.email}
        user.delete()
        log_audit(actor_id, 'DELETE', 'Users', user_info, ip_address, user_agent)

    @staticmethod
    @transaction.atomic
    def import_users_from_excel(file_obj, actor_id, ip_address=None, user_agent=None):
        try:
            wb = openpyxl.load_workbook(file_obj)
            sheet = wb.active
            created_count = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or not row[0]:
                    continue
                email, full_name, password = row[0], row[1], row[2]
                if not User.objects.filter(email=email).exists():
                    User.objects.create_user(email=email, full_name=full_name, password=password)
                    created_count += 1

            if created_count > 0:
                log_audit(actor_id, 'IMPORT_EXCEL', 'Users', {'count': created_count}, ip_address, user_agent)
            return created_count, None
        except Exception as e:
            logger.exception("Import Excel failed")
            return 0, str(e)


class AuditLogService:
    @staticmethod
    def export_to_excel(queryset):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Audit Logs"

        headers = ["ID", "User", "Action", "Module", "IP Address", "Created At"]
        ws.append(headers)

        for log in queryset:
            user_str = log.user.email if log.user else "System"
            ws.append([str(log.id), user_str, log.action, log.module, log.ip_address, str(log.created_at)])

        return wb, None

